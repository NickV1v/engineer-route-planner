from datetime import timedelta
from io import BytesIO
from zipfile import ZipFile

import pytest
from conftest import engineer, request, scenario
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from dispatch.api import create_app, with_scenario
from dispatch.baseline import solve_baseline
from dispatch.exports import build_workbook
from dispatch.geography import CoordinateEdit, GeographyStore, attach_geography
from dispatch.importer import location_id
from dispatch.models import Equipment, GeoLocation, GeoPoint
from dispatch.publication import activate_plan
from dispatch.store import PlanStore


@pytest.fixture
def saved_plan(tmp_path):
    raw = [
        request(
            "one",
            address="Ошибочный адрес",
            location_id=location_id("Ошибочный адрес"),
            service_s=1800,
            window_start_s=36000,
        ),
        request(
            "two",
            address="Второй дом",
            location_id=location_id("Второй дом"),
            skill="local",
            kind="Локальная заявка",
            service_s=900,
        ),
        request(
            "three", address="Третий дом", location_id=location_id("Третий дом"), skill="emergency"
        ),
    ]
    for index, job in enumerate(raw):
        job.source_order = index
    case = scenario(
        raw,
        [
            engineer("one", name="Анна", profile="bicycle"),
            engineer("two", name="Борис", profile="walk", skills=["local"]),
        ],
    )
    case.schedule_policy = "compact_v1"
    case.equipment_policy = "client_devices_v1"
    for job in raw:
        case.geography[job.location_id] = GeoLocation(
            location_id=job.location_id,
            address=job.address,
            status="manual",
            point=GeoPoint(
                lat=55.71,
                lon=37.61,
                label="Уточнённый дом" if job.id == "one" else job.address,
                precision="manual",
                source="dispatcher",
            ),
        )
    case.manifest["excluded_requests"] = [
        request("omitted", address="Исключённый дом").model_dump()
    ]
    case, plan = activate_plan(case, solve_baseline(case))
    case.equipment_issued = {
        "one": Equipment(routers=2, set_top_boxes=1),
        "two": Equipment(routers=0, set_top_boxes=0),
    }
    plan = plan.model_copy(update={"input_hash": case.fingerprint()})
    store = PlanStore(tmp_path / "exports.sqlite")
    response = store.create(case, with_scenario(case, {"plan": plan}))
    return case, response, store


def workbook(content):
    return load_workbook(BytesIO(content))


def rows(sheet):
    values = list(sheet.values)
    return [dict(zip(values[0], value)) for value in values[1:]]


def test_excel_order_times_metrics_names_and_exclusions(saved_plan):
    case, response, _ = saved_plan
    book = workbook(build_workbook(case, response, "excel"))
    assert book.sheetnames == ["Сводка", "Маршруты", "Инженеры", "Без назначения", "Исключённые"]
    visits = rows(book["Маршруты"])
    assert [(r["ID инженера"], r["ID заявки"]) for r in visits] == [
        ("one", "Офис"),
        ("one", "one"),
        ("two", "Офис"),
        ("two", "two"),
    ]
    job = visits[1]
    assert job["Адрес"] == "Уточнённый дом" and job["Исходный адрес"] == "Ошибочный адрес"
    assert job["Прибытие"] == timedelta(hours=10)
    assert job["Начало работы"] == timedelta(hours=10)
    assert job["Отбытие"] == timedelta(hours=10, minutes=30)
    assert visits[0]["Отбытие"] == timedelta(hours=9, minutes=50)
    assert job["В пути, мин"] == 10 and job["До точки, км"] == 1 and job["Работа, мин"] == 30
    assert job["Ожидание, мин"] == 0
    stats = rows(book["Инженеры"])
    assert stats[0]["Выдано роутеров"] == 2 and stats[0]["Выдано ТВ-приставок"] == 1
    assert stats[0]["Транспорт"] == "Велосипед"
    assert rows(book["Без назначения"])[0]["Причина"] == "Нет инженера с нужной квалификацией"
    assert rows(book["Исключённые"])[0]["ID заявки"] == "omitted"
    summary = {r[0]: r[1:] for r in book["Сводка"].values}
    assert summary["Назначено заявок"] == (2, 2)
    assert summary["Расстояние, км"] == (2, 2)
    assert summary["Без назначения"] == (None, 1)
    assert book["Маршруты"].freeze_panes == "A2"
    assert book["Маршруты"].auto_filter.ref == "A1:T5"
    assert job["Широта"] == 55.71 and job["Долгота"] == 37.61


def test_partial_export_keeps_global_rejections_and_reports_both_totals(saved_plan):
    case, response, _ = saved_plan
    book = workbook(build_workbook(case, response, "excel", ["two"]))
    assert {r["ID инженера"] for r in rows(book["Маршруты"])} == {"two"}
    assert [r["ID инженера"] for r in rows(book["Инженеры"])] == ["two"]
    summary = {r[0]: r[1:] for r in book["Сводка"].values}
    assert summary["Назначено заявок"] == (1, 2)
    assert summary["Расстояние, км"] == (1, 2)
    assert rows(book["Без назначения"])[0]["ID заявки"] == "three"
    assert rows(book["Исключённые"])[0]["ID заявки"] == "omitted"
    for ids in ([], ["two", "two"], ["unknown"]):
        with pytest.raises(ValueError):
            build_workbook(case, response, "excel", ids)


def test_printable_route_sheets_and_literal_spreadsheet_text(saved_plan):
    case, response, _ = saved_plan
    name = '=HYPERLINK("https://bad.example") /[]:*?\' ' + "И" * 120
    case.engineers[0].name = name
    case.engineers[1].name = name
    case.geography[case.requests[0].location_id].point.label = "=1+2\x01"
    book = workbook(build_workbook(case, response, "waybills"))
    assert len(book.sheetnames) == 2 and len(set(book.sheetnames)) == 2
    for sheet in book:
        assert len(sheet.title) <= 31
        assert not any(c in sheet.title for c in "[]:*?/\\")
        assert sheet.print_title_rows == "$1:$6"
        assert sheet.page_setup.orientation == "landscape"
        assert sheet.page_setup.fitToWidth == 1
        assert sheet.print_area
        assert all(cell.data_type != "f" for row in sheet for cell in row)
    assert book.worksheets[0]["C8"].value == "=1+2\n55.710000, 37.610000"
    assert book.worksheets[0]["C8"].data_type == "s"
    with ZipFile(BytesIO(build_workbook(case, response, "excel"))) as archive:
        assert all(
            b"<f>" not in archive.read(name)
            for name in archive.namelist()
            if name.startswith("xl/worksheets/")
        )
    assert len(workbook(build_workbook(case, response, "waybills", ["two"])).sheetnames) == 1


def test_empty_plan_is_exportable_without_making_up_route_sheets(saved_plan):
    case, _, _ = saved_plan
    case.engineers = []
    case.roster = None
    case.equipment_issued = None
    plan = solve_baseline(case)
    response = {"plan": plan.model_dump()}
    book = workbook(build_workbook(case, response, "excel"))
    assert rows(book["Маршруты"]) == []
    assert len(rows(book["Без назначения"])) == 3
    with pytest.raises(ValueError, match="нет назначенных"):
        build_workbook(case, response, "waybills")


def test_api_exports_pinned_snapshot_without_geocoding_routing_or_mutations(
    saved_plan, monkeypatch
):
    case, response, store = saved_plan
    app = create_app(storage_path=store.path)

    def forbidden(*args, **kwargs):
        pytest.fail("Export must only read an existing snapshot")

    monkeypatch.setattr(app.state.routing, "calculation", forbidden)
    monkeypatch.setattr(app.state.upload_geocoder, "prepare", forbidden)
    monkeypatch.setattr("dispatch.api.optimize", forbidden)
    catalog = GeographyStore(store)
    job = case.requests[0]
    catalog.edit(
        job.location_id,
        CoordinateEdit(
            address=job.address,
            expected_revision=0,
            lat=55.8,
            lon=37.8,
            selected_point=GeoPoint(
                lat=55.8,
                lon=37.8,
                label="Будущее название",
                precision="manual",
                source="dispatcher",
            ),
        ),
    )
    newer = attach_geography(case, catalog)
    newer_plan = solve_baseline(newer)
    store.commit_event(
        response["session_id"],
        1,
        "new-version",
        "hash",
        newer,
        with_scenario(newer, {"plan": newer_plan}),
    )
    with TestClient(app) as client:
        url = f"/api/sessions/{response['session_id']}/export"
        result = client.get(url, params={"version": 1, "format": "excel"})
        assert result.status_code == 200, result.text
        assert result.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert "attachment;" in result.headers["content-disposition"]
        assert rows(workbook(result.content)["Маршруты"])[1]["Адрес"] == "Уточнённый дом"
        latest = client.get(url, params={"version": 2, "format": "excel"})
        assert rows(workbook(latest.content)["Маршруты"])[1]["Адрес"] == "Будущее название"
        assert client.get(url).status_code == 422
        assert client.get(url, params={"version": 99}).status_code == 404
        assert client.get(url, params={"version": 1, "scope": "visible"}).status_code == 422
        assert client.get(url, params={"version": 1, "engineer_id": "two"}).status_code == 422
        partial = client.get(
            url,
            params={"version": 1, "format": "waybills", "scope": "visible", "engineer_id": "two"},
        )
        assert len(workbook(partial.content).sheetnames) == 1
    assert store.get(response["session_id"], 1)[0] == case
    assert store.get(response["session_id"])[1]["version"] == 2
