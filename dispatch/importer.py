"""Strictly allowlisted import: control CSV entries are never opened."""

import csv
import hashlib
import io
import json
from collections import Counter
from datetime import datetime
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

from openpyxl import load_workbook

from dispatch.models import Request

ZONES = ("Восток", "Юго-восток", "Югоцентр")
ALLOWED = {f"{zone} Синтетические данные.csv": zone for zone in ZONES}
REQUIRED_COLUMNS = {
    "Заявка",
    "Тип заявки BK",
    "Тип заявки HD",
    "Начало",
    "Окончание",
    "Район",
    "Адрес",
}
NORM_COLUMNS = [
    "Название работы",
    "Дорога до клиента/ТКД, мин.",
    "Технические работы, мин.",
    "Документы, мин.",
    "Базовый норматив, мин.",
]


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def location_id(address: str) -> str:
    # Only exact addresses share a point; do not silently geocode/merge near matches.
    return "loc-" + sha256(address.encode())[:20]


def read_norms(path: Path) -> dict:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook.active.iter_rows(values_only=True)
        headers = next(rows)
        if list(headers[:5]) != NORM_COLUMNS:
            raise ValueError("Неизвестная схема нормативов")
        result = {}
        for row in rows:
            if not any(v is not None for v in row):
                continue
            name, road, technical, documents, total = row[:5]
            if not isinstance(name, str) or name in result:
                raise ValueError("Некорректное или повторное название норматива")
            if any(
                type(v) not in (int, float) or v < 0 or not float(v).is_integer()
                for v in (road, technical, documents, total)
            ):
                raise ValueError(f"Некорректные минуты: {name}")
            if road + technical + documents != total:
                raise ValueError(f"Норматив не сходится: {name}")
            result[name] = {
                "service_s": int(technical + documents) * 60,
                "excluded_road_s": int(road) * 60,
            }
        return result
    finally:
        workbook.close()


def decoded_name(info) -> str:
    if info.flag_bits & 0x800:
        return info.filename
    try:
        return info.filename.encode("cp437").decode("cp866")
    except UnicodeEncodeError:
        # Some Python builds already honor the Info-ZIP Unicode path extra field.
        return info.filename


def parse_csv(data: bytes, filename: str, zone: str, norms: dict, policy: dict) -> dict:
    try:
        text = data.decode("utf-8-sig")
        encoding = "utf-8-sig"
    except UnicodeDecodeError:
        text = data.decode("cp1251")
        encoding = "cp1251"
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=";")
    headers = next(reader, [])
    if not REQUIRED_COLUMNS.issubset(headers) or len(set(headers)) != len(headers):
        raise ValueError(f"Неизвестные или повторные заголовки: {filename}")
    requests, errors, offices, dates = [], [], [], set()
    seen_ids = set()
    empty_rows = 0
    candidate_count = 0
    for row in reader:
        line = reader.line_num
        if not any(cell.strip() for cell in row):
            empty_rows += 1
            continue
        if row[0].strip().casefold() == "адрес офиса":
            if len(row) < 2 or not row[1].strip():
                raise ValueError(f"Пустой адрес офиса: {filename}:{line}")
            offices.append(row[1].strip())
            continue
        source_order = candidate_count
        candidate_count += 1
        try:
            if len(row) != len(headers):
                raise ValueError("Число полей не совпадает с заголовком")
            raw = dict(zip(headers, row))
            request_id = raw["Заявка"].strip()
            if not request_id or request_id in seen_ids:
                raise ValueError("Пустой или повторный ID заявки")
            start = datetime.strptime(raw["Начало"], "%d.%m.%Y %H:%M")
            end = datetime.strptime(raw["Окончание"], "%d.%m.%Y %H:%M")
            if start.date() != end.date():
                raise ValueError("Окно пересекает границу дня")
            kind = raw["Тип заявки BK"].strip()
            mapping = policy["norm_mapping"].get(kind)
            if mapping is None or mapping["work"] not in norms:
                raise ValueError(f"Нет норматива для {kind}")
            address = raw["Адрес"].strip()
            if not address:
                raise ValueError("Пустой адрес")
            request = Request(
                id=f"{zone}:{request_id}",
                zone=zone,
                source_order=source_order,
                source_line=line,
                location_id=location_id(address),
                address=address,
                kind=kind,
                subtype=raw["Тип заявки HD"].strip(),
                window_start_s=start.hour * 3600 + start.minute * 60,
                window_end_s=end.hour * 3600 + end.minute * 60,
                service_s=norms[mapping["work"]]["service_s"],
                skill=mapping["skill"],
                urgent=raw["Тип заявки HD"].strip() == "Авария",
                raw=raw,
            )
            requests.append(request)
            seen_ids.add(request_id)
            dates.add(start.date().isoformat())
        except ValueError as exc:
            errors.append({"line": line, "error": str(exc), "raw": row})
    if len(offices) != 1 or len(dates) != 1:
        raise ValueError(f"Требуется один офис и одна дата: {filename}")
    counts = Counter(r.address for r in requests)
    return {
        "id": zone,
        "date": next(iter(dates)),
        "office_address": offices[0],
        "requests": requests,
        "report": {
            "source_file": filename,
            "source_sha256": sha256(data),
            "encoding": encoding,
            "accepted": len(requests),
            "rejected": len(errors),
            "candidate_rows": candidate_count,
            "empty_rows": empty_rows,
            "errors": errors,
            "duplicate_addresses": {address: n for address, n in counts.items() if n > 1},
        },
    }


def import_sources(archive: Path, workbook: Path, policy: dict) -> list[dict]:
    norms = read_norms(workbook)
    result = []
    seen = set()
    with ZipFile(archive) as z:
        entries = z.infolist()
        if len(entries) > 100 or sum(i.file_size for i in entries) > 50 * 1024 * 1024:
            raise ValueError("Превышен лимит размера ZIP")
        for info in entries:
            name = decoded_name(info)
            path = PurePosixPath(name)
            if path.is_absolute() or ".." in path.parts or "\\" in name:
                raise ValueError("Недопустимый путь внутри ZIP")
            # Selection precedes z.read; never extract the archive wholesale.
            if path.name not in ALLOWED or info.is_dir():
                continue
            zone = ALLOWED[path.name]
            if zone in seen:
                raise ValueError(f"Повторный синтетический файл: {zone}")
            seen.add(zone)
            item = parse_csv(z.read(info), name, zone, norms, policy)
            item["norms"] = norms
            item["norms_sha256"] = sha256(workbook.read_bytes())
            result.append(item)
    if seen != set(ZONES):
        raise ValueError("В ZIP должны быть три ожидаемых синтетических CSV")
    return sorted(result, key=lambda item: ZONES.index(item["id"]))


def read_policy(path: Path) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))
